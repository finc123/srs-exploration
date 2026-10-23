import pandas as pd, json
from openpyxl.utils import get_column_letter
ORDER = ['New York City','Paris','London','Rome','Barcelona','Tokyo','Mexico City','Madrid','Amsterdam','San Francisco','Chicago',
 'Las Vegas','Orlando','Miami Beach','Milan','Lisbon','Los Angeles','Miami','New Orleans','San Diego','Philadelphia','Cancún','Toronto',
 'Singapore','Washington',
 'Berlin','Florence','Atlanta','Houston','Kissimmee','Austin','Seattle','Boston','Nashville','Kyoto','Seoul','Bangkok','Sydney',
 'San Juan','Montreal','Vancouver','Venice','Prague','Vienna','Budapest','Athens','Dublin','Istanbul','Dubai','Osaka','Tulum','Melbourne',
 'Charlotte','Denver','Dallas','Edinburgh','Nice','Brussels','Jacksonville','Portland','Newark','Fort Lauderdale','Cambridge','Baltimore',
 'San Jose','Pittsburgh','Honolulu','Buenos Aires','Rio de Janeiro',
 'Munich','Porto','Seville','Naples','Valencia','Málaga','Copenhagen','Zurich','Lyon','Montpellier','Marseille','Stockholm','Hong Kong',
 'Bali','Cape Town','Marrakesh','Playa del Carmen','Medellín','Mumbai','New Delhi']
p=pd.read_csv('place_ids.csv',encoding='utf-8')
assert len(ORDER)==len(set(ORDER))==len(p), (len(ORDER),len(p))
sup=pd.DataFrame(json.load(open('supply.json',encoding='utf-8')))
sup.columns=[c.lower() for c in sup.columns]
p=p.merge(sup[['city','hotels_15km','reviews_15km']],on='city',how='left').fillna({'hotels_15km':0,'reviews_15km':0})
t=pd.read_csv(r'C:\Users\fintan.creedon\code\exploration\abnb\hotels\output\top_cities\top_cities_hotel_or_room.csv')
us=t[t.country=='US'][['state','city','confirmations','conf_hotel_or_room']].copy()
us['us_rank']=range(1,len(us)+1)
p=p.merge(us,on=['state','city'],how='left')
cty=pd.read_csv(r'C:\Users\fintan.creedon\code\exploration\abnb\hotels\output\top_cities\countries.csv')
CMAP={'Canada':'CA','United Kingdom':'(n/a)','Singapore':'(n/a)','South Africa':'(n/a)','Türkiye':'TURKEY','Czechia':'CZECH REPUBLIC','South Korea':'SOUTH KOREA','United Arab Emirates':'UNITED ARAB EMIRATES','Hong Kong':'HONG KONG'}
cty['crank']=range(1,len(cty)+1)
p['ckey']=p.country.map(lambda c: CMAP.get(c,c.upper()))
p=p.merge(cty[['country','confirmations','crank']].rename(columns={'country':'ckey','confirmations':'country_conf'}),on='ckey',how='left')
p['priority']=p.city.map({c:i+1 for i,c in enumerate(ORDER)})
p=p.sort_values('priority')
p['tier']=pd.cut(p.priority,[0,25,52,99],labels=['1 - scrape first','2','3'])
def why(r):
    s=[]
    if r.country=='United States':
        s.append(f"US #{int(r.us_rank)} by Hotel room + Room stays in panel" if r.us_rank<=30 else "Hawaii absent from current scrape")
    else:
        if pd.notna(r.country_conf): s.append(f"{int(r.country_conf):,} panel stays in country")
        else: s.append("country not reliably parsed in NIQ addresses")
    s.append('not in current scrape' if r.hotels_15km==0 else f"{int(r.hotels_15km)} hotels found in Oct-05 scan")
    return '; '.join(s)
p['rationale']=p.apply(why,axis=1)
out=pd.DataFrame({'Priority':p.priority,'Tier':p.tier.astype(str),'City':p.city,'State':p.state.fillna(''),'Country':p.country,
 'Suggested Airbnb query':p.vendor_query.fillna(''),'Place ID (from 72degree lists)':p.place_id.fillna('to resolve'),
 'Hotels in Oct-05 scan (15 km)':p.hotels_15km.fillna(0).astype(int),'Reviews on those hotels':p.reviews_15km.fillna(0).astype(int),
 'Panel Hotel room + Room stays (US)':p.conf_hotel_or_room,'Panel stays, all types (US city)':p.confirmations_x if 'confirmations_x' in p else p.confirmations,
 'Panel stays in country':p.country_conf,'Rationale':p.rationale})
f=r'C:\Users\fintan.creedon\code\exploration\abnb\docs\abnb_scrape_city_list.xlsx'
with pd.ExcelWriter(f,engine='openpyxl') as w:
    out.to_excel(w,sheet_name='Scrape cities',index=False)
    ws=w.sheets['Scrape cities']; ws.freeze_panes='D2'; ws.auto_filter.ref=ws.dimensions
    for i,col in enumerate(out.columns,1):
        L=get_column_letter(i)
        ws.column_dimensions[L].width={'Rationale':60,'Suggested Airbnb query':42,'Place ID (from 72degree lists)':32,'City':18,'Country':20,'Tier':14}.get(col,max(10,len(col)*0.8))
        if out[col].dtype.kind in 'if' and col!='Priority':
            for c in ws[L][1:]: c.number_format='#,##0'
print(out.drop(columns=['Rationale','Suggested Airbnb query']).to_string(index=False))
